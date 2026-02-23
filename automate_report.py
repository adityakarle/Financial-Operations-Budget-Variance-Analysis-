import pandas as pd
import psycopg2
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from datetime import datetime
import os

DB_CONFIG = {
    'host': 'localhost', 'database': 'budget_analysis',
    'user': 'postgres', 'password': 'Admin', 'port': 5432
}

conn = psycopg2.connect(**DB_CONFIG)

variance_df = pd.read_sql_query('''
    WITH actual AS (
        SELECT d.department_name,
               DATE_TRUNC('month', t.date) AS month,
               ROUND(SUM(t.amount)::numeric, 2) AS actual_spend
        FROM transactions t
        JOIN departments d ON t.department_id = d.department_id
        WHERE t.status = 'Approved'
        GROUP BY d.department_name, DATE_TRUNC('month', t.date)
    ),
    budget AS (
        SELECT d.department_name, b.month,
               ROUND(SUM(b.budgeted_amount)::numeric, 2) AS total_budget
        FROM budgets b
        JOIN departments d ON b.department_id = d.department_id
        GROUP BY d.department_name, b.month
    )
    SELECT a.department_name, a.month, a.actual_spend, b.total_budget,
           ROUND((a.actual_spend - b.total_budget)::numeric, 2) AS variance,
           ROUND(((a.actual_spend - b.total_budget) / NULLIF(b.total_budget,0) * 100)::numeric,1) AS variance_pct
    FROM actual a
    JOIN budget b ON a.department_name = b.department_name AND a.month = b.month
    ORDER BY variance_pct
''', conn)

category_df = pd.read_sql_query('''
    SELECT c.category_name,
           COUNT(*) AS transactions,
           ROUND(SUM(t.amount)::numeric, 2) AS total_spend,
           ROUND(AVG(t.amount)::numeric, 2) AS avg_amount
    FROM transactions t
    JOIN categories c ON t.category_id = c.category_id
    WHERE t.status = 'Approved'
    GROUP BY c.category_name
    ORDER BY total_spend DESC
''', conn)

conn.close()

wb = openpyxl.Workbook()

# --- Sheet 1: Variance Analysis ---
ws1 = wb.active
ws1.title = 'Monthly Variance'

headers = ['Department', 'Month', 'Actual Spend', 'Budget', 'Variance', 'Variance %']
ws1.append(headers)

# Style header row
for cell in ws1[1]:
    cell.fill = PatternFill('solid', fgColor='1F4E79')
    cell.font = Font(bold=True, color='FFFFFF', size=11)
    cell.alignment = Alignment(horizontal='center')

# Add data with color coding
red   = PatternFill('solid', fgColor='FFCCCC')
green = PatternFill('solid', fgColor='CCFFCC')

for _, row in variance_df.iterrows():
    row['month'] = row['month'].replace(tzinfo=None)
    ws1.append(list(row))
    last_row = ws1.max_row
    pct = float(row['variance_pct'])
    if abs(pct) >= 15:
        for cell in ws1[last_row]:
            cell.fill = red
    elif abs(pct) <= 3:
        for cell in ws1[last_row]:
            cell.fill = green

# Auto width
for col in ws1.columns:
    max_len = max(len(str(c.value or '')) for c in col) + 4
    ws1.column_dimensions[get_column_letter(col[0].column)].width = min(max_len, 30)

# --- Sheet 2: Category Summary ---
ws2 = wb.create_sheet('Category Summary')
ws2.append(['Category', 'Transactions', 'Total Spend', 'Avg Amount'])

for cell in ws2[1]:
    cell.fill = PatternFill('solid', fgColor='1F4E79')
    cell.font = Font(bold=True, color='FFFFFF', size=11)
    cell.alignment = Alignment(horizontal='center')

for _, row in category_df.iterrows():
    ws2.append(list(row))

for col in ws2.columns:
    max_len = max(len(str(c.value or '')) for c in col) + 4
    ws2.column_dimensions[get_column_letter(col[0].column)].width = min(max_len, 30)

# Save
os.makedirs('../reports', exist_ok=True)
filename = f'../reports/financial_report_{datetime.now().strftime("%Y%m")}.xlsx'
wb.save(filename)
print(f'Report saved: {filename}')
